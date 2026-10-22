import io
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

import httpx
from openpyxl import Workbook

from seed_data import seed_if_empty
from server import _download_workbook, create_app, hash_password, verify_password


def participant(nip: str, name: str) -> dict:
    return {
        "sourceRow": 2,
        "nip": nip,
        "nama": name,
        "jab": "Analis SDMA",
        "jt": "Analis SDM Aparatur Ahli Madya",
        "jenjang": "Ahli Madya",
        "inst": "Instansi Uji",
        "nilai": 4,
        "ld": 4,
        "hasil": "Memenuhi",
        "batch": "Batch 1",
        "mulai": "2026-09-03T14:11:43",
        "sel": "2026-09-03T14:33:33",
        "dup": 1,
        "chk": "OK",
        "s": [4, 3],
    }


class WorkbookDownloadTests(unittest.TestCase):
    def test_downloads_once_and_reuses_cached_workbook(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            destination = Path(temp_dir) / "seed.xlsx"
            with patch("server.urlopen", return_value=io.BytesIO(b"xlsx data")) as download:
                result = _download_workbook("https://example.test/seed.xlsx", destination)
                cached_result = _download_workbook("https://example.test/changed.xlsx", destination)

            self.assertEqual(result, destination)
            self.assertEqual(cached_result, destination)
            self.assertEqual(destination.read_bytes(), b"xlsx data")
            download.assert_called_once()

    def test_rejects_invalid_url_and_empty_download(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            destination = Path(temp_dir) / "seed.xlsx"
            with self.assertRaises(ValueError):
                _download_workbook("file:///seed.xlsx", destination)

            with patch("server.urlopen", return_value=io.BytesIO()):
                with self.assertRaisesRegex(ValueError, "kosong"):
                    _download_workbook("https://example.test/seed.xlsx", destination)
            self.assertFalse(destination.exists())


class AuthenticationTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.db_path = Path(self.temp_dir.name) / "test.sqlite3"
        self.workbook_path = Path(self.temp_dir.name) / "seed.xlsx"
        self.unit_names_workbook_path = Path(self.temp_dir.name) / "unit-names.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "NIP", "Nama", "Jabatan", "Jenjang Jabatan", "Jenjang", "INSTANSI",
            "Nilai", "Level Dominan", "Hasil", "Batch", "Mulai", "Selesai",
            "DUPLIKAT", "Check Level", "Soal 1", "Soal 2",
        ])
        sheet.append([
            "200000000000000001", "Seed Alya", "Analis SDMA",
            "Analis SDM Aparatur Ahli Madya", "Ahli Madya", "Instansi Seed",
            4.25, 4, "Memenuhi", "Batch Seed", datetime(2026, 9, 3, 14, 11, 43),
            datetime(2026, 9, 3, 14, 33, 33), 1, "OK", 4, 3,
        ])
        sheet.append([
            "200000000000000002", "Seed Budi", "Analis SDMA",
            "Analis SDM Aparatur Ahli Madya", "Ahli Madya", "Instansi Seed",
            3, 3, "Belum Memenuhi", "Batch Seed", datetime(2026, 9, 3, 14, 11, 43),
            datetime(2026, 9, 3, 14, 33, 33), 1, "OK", 3, 3,
        ])
        workbook.save(self.workbook_path)
        unit_names_workbook = Workbook()
        unit_names_sheet = unit_names_workbook.active
        unit_names_sheet.append([
            "No", "Soal", "yang dimasukan ke nama Unit kompetensi / Nomor Soal",
            None, None, None,
            "No", "Soal", "yang dimasukan ke nama Unit kompetensi / Nomor Soal",
            None, None, None,
            "No", "Soal", "yang dimasukan ke nama Unit kompetensi / Nomor Soal",
        ])
        unit_names_sheet.append([
            1, "Kompetensi teknis umum", "Unit Analis",
            None, None, None,
            1, "Kompetensi teknis umum", "Unit Auditor",
            None, None, None,
            1, "Kompetensi teknis umum", "Unit Asesor",
        ])
        unit_names_workbook.save(self.unit_names_workbook_path)
        self.app = create_app(
            db_path=self.db_path,
            admin_password="test-admin-password-123",
            session_secret="test-session-secret-value-at-least-32-bytes",
            cookie_secure=False,
            seed_workbook_path=self.workbook_path,
            unit_names_workbook_path=self.unit_names_workbook_path,
        )
        self.lifespan = self.app.router.lifespan_context(self.app)
        await self.lifespan.__aenter__()
        self.client = httpx.AsyncClient(
            transport=httpx.ASGITransport(app=self.app),
            base_url="http://testserver",
        )

    async def asyncTearDown(self):
        await self.client.aclose()
        await self.lifespan.__aexit__(None, None, None)
        self.temp_dir.cleanup()

    async def login_admin_and_import(self):
        response = await self.client.post(
            "/api/auth/login",
            json={"username": "admin", "password": "test-admin-password-123"},
        )
        self.assertEqual(response.status_code, 200)
        response = await self.client.post(
            "/api/admin/import",
            json={
                "file": "test.xlsx",
                "soalCount": 2,
                "rows": [
                    participant("100000000000000001", "Álya Peserta"),
                    participant("100000000000000002", "Budi Peserta"),
                ],
            },
        )
        self.assertEqual(response.status_code, 200)

    def test_password_hash_is_salted_and_verified(self):
        encoded = hash_password("a-private-password")
        self.assertTrue(verify_password("a-private-password", encoded))
        self.assertFalse(verify_password("wrong-password", encoded))
        self.assertNotEqual(encoded, hash_password("a-private-password"))

    async def test_startup_seeds_dataset_once_and_creates_initial_nip_passwords(self):
        response = await self.client.post(
            "/api/auth/login",
            json={"username": "admin", "password": "test-admin-password-123"},
        )
        self.assertEqual(response.status_code, 200)
        data = await self.client.get("/api/data")
        self.assertEqual(data.status_code, 200)
        self.assertEqual(data.json()["file"], "seed.xlsx")
        self.assertEqual(data.json()["soalCount"], 2)
        self.assertEqual(len(data.json()["rows"]), 2)
        self.assertEqual(data.json()["rows"][0]["nilai"], 4.25)
        self.assertEqual(data.json()["rows"][0]["s"], [4, 3])
        self.assertEqual(data.json()["unitNames"]["Analis"]["1"], "Unit Analis")
        self.assertEqual(data.json()["unitNames"]["Auditor"]["1"], "Unit Auditor")
        self.assertEqual(data.json()["unitNames"]["Asesor"]["1"], "Unit Asesor")

        await self.client.post("/api/auth/logout")
        participant_login = await self.client.post(
            "/api/auth/login",
            json={"username": "sEEd aLYa", "password": "200000000000000001"},
        )
        self.assertEqual(participant_login.status_code, 200)
        self.assertTrue(participant_login.json()["must_change_password"])
        missing_workbook = Path(self.temp_dir.name) / "missing.xlsx"
        self.assertFalse(seed_if_empty(self.db_path, missing_workbook))

        admin = await self.client.post(
            "/api/auth/login",
            json={"username": "admin", "password": "test-admin-password-123"},
        )
        self.assertEqual(admin.status_code, 200)
        unchanged = await self.client.get("/api/data")
        self.assertEqual(unchanged.json()["file"], "seed.xlsx")
        self.assertEqual(len(unchanged.json()["rows"]), 2)

    async def test_admin_sees_all_rows_and_user_is_restricted_to_own_row(self):
        await self.login_admin_and_import()
        admin_data = await self.client.get("/api/data")
        self.assertEqual(admin_data.status_code, 200)
        self.assertEqual(len(admin_data.json()["rows"]), 2)

        await self.client.post("/api/auth/logout")
        response = await self.client.post(
            "/api/auth/login",
            json={"username": "áLYA PESERTA", "password": "100000000000000001"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["must_change_password"])
        self.assertEqual((await self.client.get("/api/data")).status_code, 403)
        other_client = httpx.AsyncClient(
            transport=httpx.ASGITransport(app=self.app),
            base_url="http://testserver",
        )
        other_login = await other_client.post(
            "/api/auth/login",
            json={"username": "Álya Peserta", "password": "100000000000000001"},
        )
        self.assertEqual(other_login.status_code, 200)

        response = await self.client.post(
            "/api/auth/change-password",
            json={
                "current_password": "100000000000000001",
                "new_password": "new-private-password-123",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual((await other_client.get("/api/data")).status_code, 401)
        await other_client.aclose()
        forbidden_import = await self.client.post(
            "/api/admin/import",
            json={
                "file": "forbidden.xlsx",
                "soalCount": 2,
                "rows": [participant("100000000000000001", "Álya Peserta")],
            },
        )
        self.assertEqual(forbidden_import.status_code, 403)
        data = await self.client.get("/api/data")
        self.assertEqual(data.status_code, 200)
        self.assertEqual([row["nip"] for row in data.json()["rows"]], ["100000000000000001"])
        auth = (await self.client.get("/api/auth/me")).json()
        self.assertEqual(auth["role"], "user")
        self.assertEqual(auth["nip"], "100000000000000001")

    async def test_import_preserves_changed_password_and_rejects_cross_origin(self):
        await self.login_admin_and_import()
        await self.client.post("/api/auth/logout")
        await self.client.post(
            "/api/auth/login",
            json={"username": "Álya Peserta", "password": "100000000000000001"},
        )
        changed = await self.client.post(
            "/api/auth/change-password",
            json={
                "current_password": "100000000000000001",
                "new_password": "new-private-password-123",
            },
        )
        self.assertEqual(changed.status_code, 200)
        await self.client.post("/api/auth/logout")
        await self.login_admin_and_import()
        await self.client.post("/api/auth/logout")
        self.assertEqual(
            (await self.client.post(
                "/api/auth/login",
                json={"username": "Álya Peserta", "password": "new-private-password-123"},
            )).status_code,
            200,
        )
        blocked = await self.client.post(
            "/api/auth/logout",
            headers={"Origin": "https://attacker.example"},
        )
        self.assertEqual(blocked.status_code, 403)

    async def test_dashboard_is_served_without_embedded_dataset(self):
        response = await self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("const EMBEDDED = null;", response.text)
        self.assertNotIn("100000000000000001", response.text)
        self.assertEqual((await self.client.get("/api/auth/me")).json(), {"authenticated": False})
        self.assertEqual((await self.client.get("/api/data")).status_code, 401)
        self.assertEqual((await self.client.get("/index.html")).status_code, 404)
        self.assertEqual((await self.client.get("/PRISMA_DATA_FINAL.xlsx")).status_code, 404)


if __name__ == "__main__":
    unittest.main()
