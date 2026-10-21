from __future__ import annotations

import json
import math
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from dataset_store import connect_db, persist_dataset, validate_dataset

ROOT = Path(__file__).resolve().parent
DEFAULT_WORKBOOK = ROOT / "PRISMA_DATA_FINAL.xlsx"
DEFAULT_UNIT_NAMES_WORKBOOK = ROOT / "NAMA UNIT KOMPETENSI.xlsx"
UNIT_NAME_COLUMNS = ((0, "Analis"), (6, "Auditor"), (12, "Asesor"))
MAX_ROWS = 50_000
MAX_SOAL_COUNT = 200


def _cell_string(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _cell_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    match = re.match(r"^\s*[+-]?(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?", str(value))
    if not match:
        return None
    number = float(match.group(0))
    return number if math.isfinite(number) else None


def _js_round(value: float) -> int:
    return math.floor(value + 0.5)


def _iso_datetime(value: Any) -> str | None:
    if not isinstance(value, (datetime, date)):
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%dT%H:%M:%S")
    return datetime.combine(value, datetime.min.time()).strftime("%Y-%m-%dT%H:%M:%S")


def read_workbook(path: Path = DEFAULT_WORKBOOK) -> dict[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(
            f"File seed Excel tidak ditemukan: {path}. Letakkan PRISMA_DATA_FINAL.xlsx di folder proyek."
        )

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        header_row = None
        header_index = None
        for index, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            normalized = [_cell_string(value) for value in values]
            if "NIP" in normalized and "Nama" in normalized:
                header_row = normalized
                header_index = index
                break
        if header_row is None or header_index is None:
            raise ValueError("Baris header (NIP, Nama) tidak ditemukan di sheet pertama.")

        columns: dict[str, int] = {}
        for index, name in enumerate(header_row):
            if name and name not in columns:
                columns[name] = index

        required = ("NIP", "Nama", "Jabatan", "Jenjang Jabatan", "Nilai", "Hasil")
        missing = [name for name in required if name not in columns]
        if missing:
            raise ValueError("Kolom wajib tidak ada: " + ", ".join(missing))

        score_columns = []
        score_number = 1
        while f"Soal {score_number}" in columns:
            score_columns.append(columns[f"Soal {score_number}"])
            score_number += 1
        if not score_columns or len(score_columns) > MAX_SOAL_COUNT:
            raise ValueError(f"Jumlah kolom soal harus antara 1 dan {MAX_SOAL_COUNT}.")

        check_key = next((key for key in columns if re.match(r"^Check", key, re.IGNORECASE)), None)

        def get(values: tuple[Any, ...], key: str) -> Any:
            column = columns.get(key)
            return values[column] if column is not None and column < len(values) else None

        def get_at(values: tuple[Any, ...], column: int) -> Any:
            return values[column] if column < len(values) else None

        rows = []
        for source_row, values in enumerate(
            sheet.iter_rows(min_row=header_index + 1, values_only=True),
            start=header_index + 1,
        ):
            nip = _cell_string(get(values, "NIP"))
            name = _cell_string(get(values, "Nama"))
            if not nip and not name:
                continue

            score_value = _cell_number(get(values, "Nilai"))
            dominant_level = _cell_number(get(values, "Level Dominan"))
            duplicate = _cell_number(get(values, "DUPLIKAT"))
            scores = [
                None if (number := _cell_number(get_at(values, column))) is None else _js_round(number)
                for column in score_columns
            ]
            row = {
                "sourceRow": source_row,
                "nip": nip,
                "nama": name,
                "jab": _cell_string(get(values, "Jabatan")),
                "jt": _cell_string(get(values, "Jenjang Jabatan")),
                "jenjang": _cell_string(get(values, "Jenjang")),
                "inst": _cell_string(get(values, "INSTANSI")),
                "nilai": None if score_value is None else _js_round(score_value * 100) / 100,
                "ld": None if dominant_level is None else _js_round(dominant_level),
                "hasil": _cell_string(get(values, "Hasil")),
                "batch": _cell_string(get(values, "Batch")),
                "mulai": _iso_datetime(get(values, "Mulai")),
                "sel": _iso_datetime(get(values, "Selesai")),
                "dup": None if duplicate is None else _js_round(duplicate),
                "chk": _cell_string(get(values, check_key)) if check_key else "",
                "s": scores,
            }
            rows.append(row)
            if len(rows) > MAX_ROWS:
                raise ValueError(f"Jumlah baris peserta melebihi batas {MAX_ROWS}.")

        if not rows:
            raise ValueError("Tidak ada baris data peserta pada sheet pertama.")

        return {"file": path.name, "soalCount": len(score_columns), "rows": rows}
    finally:
        workbook.close()


def read_unit_names(path: Path = DEFAULT_UNIT_NAMES_WORKBOOK) -> dict[str, dict[str, str]]:
    if not path.is_file():
        raise FileNotFoundError(f"File nama unit kompetensi tidak ditemukan: {path}.")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        unit_names: dict[str, dict[str, str]] = {}

        for start, passport in UNIT_NAME_COLUMNS:
            headers = [_cell_string(value).casefold() for value in header[start : start + 3]]
            if (
                len(headers) != 3
                or headers[0] != "no"
                or headers[1] != "soal"
                or not headers[2].startswith("yang dimasukan ke nama unit kompetensi")
            ):
                raise ValueError(
                    f"Kolom daftar unit kompetensi untuk {passport} tidak sesuai format."
                )

            names: dict[str, str] = {}
            for values in sheet.iter_rows(min_row=2, min_col=start + 1, max_col=start + 3, values_only=True):
                number_value, _, name_value = values
                if number_value is None and name_value is None:
                    continue
                number = _cell_number(number_value)
                name = " ".join(_cell_string(name_value).split())
                if number is None or number <= 0 or not number.is_integer() or not name:
                    raise ValueError(
                        f"Nomor atau nama unit kompetensi {passport} tidak valid."
                    )
                number_key = str(int(number))
                if number_key in names:
                    raise ValueError(
                        f"Nomor unit kompetensi {number_key} berulang untuk {passport}."
                    )
                names[number_key] = name
            if not names:
                raise ValueError(f"Daftar unit kompetensi {passport} kosong.")
            unit_names[passport] = names

        return unit_names
    finally:
        workbook.close()


def save_unit_names(db_path: Path, unit_names: dict[str, dict[str, str]]) -> None:
    with connect_db(db_path) as connection:
        connection.execute(
            """
            INSERT INTO metadata (key, value) VALUES ('unitNames', ?)
            ON CONFLICT(key) DO UPDATE SET value = excluded.value
            """,
            (json.dumps(unit_names, ensure_ascii=False, separators=(",", ":")),),
        )


def seed_if_empty(db_path: Path, workbook_path: Path = DEFAULT_WORKBOOK) -> bool:
    with connect_db(db_path) as connection:
        row_count = connection.execute("SELECT COUNT(*) FROM dataset_rows").fetchone()[0]
    if row_count:
        return False

    data = read_workbook(workbook_path)
    rows, participants = validate_dataset(data["rows"], data["soalCount"])

    with connect_db(db_path) as connection:
        connection.execute("BEGIN IMMEDIATE")
        row_count = connection.execute("SELECT COUNT(*) FROM dataset_rows").fetchone()[0]
        if row_count:
            return False
        persist_dataset(connection, data["file"], data["soalCount"], rows, participants)
    return True
